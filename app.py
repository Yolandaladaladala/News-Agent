import io, json, os, re, time, html, hashlib
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import urlparse
import pandas as pd
import requests
import streamlit as st
from bs4 import BeautifulSoup
from dateutil import parser as dtparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE=Path(__file__).parent
RULES=json.loads((BASE/'searching_rules/config.json').read_text(encoding='utf-8'))
REGIONS=RULES['regions']
COLS=['region_code','region_name_zh','title','date','location','tags','summary','url','url_label','icon_img_url']
AUDIT=['score','status','evidence','source_date','source_domain','source_access','reason','selected','order']
st.set_page_config(page_title='XSTAR News Research Agent',layout='wide',page_icon='📰')
st.title('XSTAR · 全球汽车行业新闻研究 Agent')
st.caption('日期 / 可选线索 → 多语言搜索 → 原文核验 → AI评分 → 编辑 → Excel + Outlook HTML')

def secret(name,default=''):
    try: return st.secrets.get(name,os.getenv(name,default))
    except Exception: return os.getenv(name,default)

def serper(query,start,end,key,num=10):
    r=requests.post('https://google.serper.dev/search',headers={'X-API-KEY':key,'Content-Type':'application/json'},json={'q':f'{query} after:{start.isoformat()} before:{(end+timedelta(days=1)).isoformat()}','num':num},timeout=25)
    r.raise_for_status()
    return [{'title':x.get('title',''),'url':x.get('link',''),'snippet':x.get('snippet',''),'date_hint':x.get('date','')} for x in r.json().get('organic',[])]

def scrape(url):
    try:
        r=requests.get(url,timeout=14,headers={'User-Agent':'Mozilla/5.0 (compatible; XSTARResearch/1.0)'},allow_redirects=True)
        if r.status_code!=200:return {'text':'','date':'','access':f'HTTP {r.status_code}'}
        soup=BeautifulSoup(r.text,'html.parser')
        for x in soup(['script','style','nav','footer','header']):x.decompose()
        meta=[]
        for attr,key in [('property','article:published_time'),('name','date'),('itemprop','datePublished'),('name','pubdate'),('name','datePublished')]:
            tag=soup.find('meta',attrs={attr:key})
            if tag and tag.get('content'):meta.append(tag['content'])
        t=soup.find('time',attrs={'datetime':True})
        if t:meta.append(t['datetime'])
        for node in soup.select('article p, main p')[:80]:pass
        paras=[x.get_text(' ',strip=True) for x in soup.select('article p, main p') if len(x.get_text(' ',strip=True))>35]
        if len(paras)<3:paras=[x.get_text(' ',strip=True) for x in soup.find_all('p') if len(x.get_text(' ',strip=True))>35]
        text=' '.join(paras)[:15000]
        blocked=bool(re.search(r'subscribe to (read|continue)|sign in to (read|continue)|already a subscriber|register to read',text[:3000],re.I))
        return {'text':text,'date':meta[0] if meta else '', 'access':'paywall_possible' if blocked else ('accessible' if len(text)>350 else 'insufficient_text')}
    except Exception as e:return {'text':'','date':'','access':f'error: {type(e).__name__}'}

def parse_date(v):
    if not v:return ''
    try:return dtparse.parse(str(v),fuzzy=True).date().isoformat()
    except Exception:return ''

def ai_json(prompt,key,model):
    r=requests.post('https://openrouter.ai/api/v1/chat/completions',headers={'Authorization':f'Bearer {key}','Content-Type':'application/json'},json={'model':model,'temperature':0,'response_format':{'type':'json_object'},'messages':[{'role':'system','content':'你是严格的新闻事实核验研究员。只使用给定原文证据，不得编造。只输出JSON。'},{'role':'user','content':prompt}]},timeout=85)
    r.raise_for_status()
    s=r.json()['choices'][0]['message']['content']
    try:return json.loads(s)
    except Exception:
        m=re.search(r'\{.*\}',s,re.S)
        return json.loads(m.group(0)) if m else {}

def country_region(s):
    s=str(s).lower()
    rules=[('北美地区','美国 usa united states 加拿大 canada 墨西哥 mexico'),('拉美地区','巴西 brazil 阿根廷 argentina 智利 chile 哥伦比亚 colombia'),('欧洲地区','英国 uk britain 德国 germany 法国 france 欧洲 europe 欧盟 eu 意大利 italy'),('中东地区','中东 middle east 沙特 saudi 阿联酋 uae 伊朗 iran 以色列 israel'),('非洲地区','非洲 africa 南非 south africa 埃及 egypt')]
    for reg,words in rules:
        if any(w in s for w in words.split(' ')):return reg
    return '亚太地区'

def extract_clues(file,txt):
    clues=[x.strip() for x in txt.splitlines() if x.strip()]
    old=[]
    if file:
        w=pd.ExcelFile(file)
        sh=next((s for s in w.sheet_names if 'news' in s.lower() or '新闻' in s),w.sheet_names[-1])
        df=pd.read_excel(file,sheet_name=sh,usecols=lambda c:str(c) in COLS)
        for _,r in df.iterrows():
            title=str(r.get('title','')).strip()
            if title and title.lower()!='nan':clues.append(title);old.append(title)
    return clues[:35],old

def research(start,end,clues,serper_key,ai_key,model,progress):
    queries=[]
    for q in RULES['topics']:queries.append(q)
    for country,qs in RULES['local_queries'].items():
        queries.extend(qs)
    for clue in clues:queries.append(clue[:110])
    queries=queries[:RULES['max_queries']]
    seen=set(); candidates=[]; errors=[]
    for i,q in enumerate(queries):
        progress.progress((i+1)/max(1,len(queries))*0.35,text=f'多语言搜索 {i+1}/{len(queries)} · {q[:38]}')
        try:
            for x in serper(q,start,end,serper_key,RULES['results_per_query']):
                key=x['url'].split('?')[0].rstrip('/')
                if key and key not in seen:seen.add(key);candidates.append(x)
        except Exception as e:errors.append(f'搜索失败 {q[:40]}: {e}')
    if not candidates:raise RuntimeError('没有获得任何搜索结果。请检查 Serper Key、额度及网络。')
    # Search rank can be noisy; prioritize accessible official sources without excluding other credible media.
    def rank(x):
        d=urlparse(x['url']).netloc.lower()
        return (0 if any(d.endswith(z) for z in RULES['primary_domains']) else 1, -sum(k in (x['title']+' '+x['snippet']).lower() for k in ['loan','credit','bank','auto','vehicle','car','finance','สินเชื่อ','汽车','รถยนต์']))
    candidates.sort(key=rank)
    rows=[]
    limit=min(len(candidates),RULES['max_articles_to_review'])
    for i,x in enumerate(candidates[:limit]):
        progress.progress(.35+.65*(i+1)/max(1,limit),text=f'访问原文并分析 {i+1}/{limit}')
        doc=scrape(x['url']); d=parse_date(doc['date'])
        # Search snippet date is not sufficiently authoritative for auto-selection.
        valid_date=bool(d and start.isoformat()<=d<=end.isoformat())
        evidence=doc['text'][:10500]
        if len(evidence)<350 or not valid_date or doc['access']!='accessible':
            rows.append({'region_code':country_region(x['title']),'region_name_zh':'','title':x['title'],'date':d,'location':'','tags':'','summary':x['snippet'],'url':x['url'],'url_label':urlparse(x['url']).netloc,'icon_img_url':'','score':0,'status':'待核验','evidence':'','source_date':d,'source_domain':urlparse(x['url']).netloc,'source_access':doc['access'],'reason':'原文日期缺失/超范围或正文不可访问','selected':False,'order':i+1})
            continue
        prompt=f'''根据下方原文判断新闻是否值得进入XSTAR全球汽车金融行业快讯。研究日期 {start} 至 {end}。必须逐字摘取原文中支持核心事件及数字的证据片段（连续20-160字符）。不得使用片段之外的信息。\n评分权重：业务相关性30，行业影响25，来源可信度20，时效性15，可行动性10。\n返回JSON对象字段 title(中文), location(中文国家), region(六区之一), tags(2-3个中文标签用；分隔), summary(中文70-150字), evidence(原文逐字证据), score(0-100整数), reason(简短), supported(boolean)。\n标题:{x['title']}\n来源:{x['url']}\n正文:{evidence}'''
        try:
            a=ai_json(prompt,ai_key,model)
            quote=str(a.get('evidence','')).strip()
            supported=bool(a.get('supported')) and len(quote)>=20 and quote in evidence
            score=max(0,min(100,int(a.get('score',0)))) if supported else 0
            region=a.get('region') if a.get('region') in REGIONS else country_region(a.get('location','')+' '+x['title'])
            rows.append({'region_code':region,'region_name_zh':str(a.get('location','')),'title':str(a.get('title') or x['title']),'date':d,'location':str(a.get('location','')),'tags':str(a.get('tags','')),'summary':str(a.get('summary','')),'url':x['url'],'url_label':urlparse(x['url']).netloc,'icon_img_url':'','score':score,'status':'已核验' if supported else '待核验','evidence':quote if supported else '', 'source_date':d,'source_domain':urlparse(x['url']).netloc,'source_access':doc['access'],'reason':str(a.get('reason','')) if supported else 'AI证据无法在原文定位','selected':supported and score>=RULES['min_score'],'order':i+1})
        except Exception as e:errors.append(f'AI分析失败 {x["title"][:35]}: {e}')
    # Near-title duplicate filter; preserves candidates for review.
    tokens=set()
    for row in sorted(rows,key=lambda r:-r['score']):
        key=re.sub(r'[^\w\u4e00-\u9fff]','',row['title'].lower())[:48]
        if key in tokens:row['selected']=False;row['status']='疑似重复';row['reason']='标题重复'
        tokens.add(key)
    chosen=sorted((r for r in rows if r['selected']),key=lambda r:-r['score'])
    for r in chosen[RULES['max_report_items']:]:r['selected']=False
    return rows,errors,len(candidates)

def excel_bytes(df,start,end,issue):
    w=Workbook();m=w.active;m.title='meta';m.append(['key','value'])
    for k,v in [('report_title','全球汽车行业快讯'),('report_date',f'{end.year}年{end.month}月'),('report_issue',issue),('preheader','全球汽车金融、行业及宏观重要动态'),('last_updated',str(date.today()))]:m.append([k,v])
    s=w.create_sheet('news');s.append(COLS)
    for _,r in df.iterrows():s.append([str(r.get(k,'') or '') for k in COLS])
    for sheet in (m,s):
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for c in sheet[1]:c.fill=PatternFill('solid',fgColor='16233B');c.font=Font(bold=True,color='FFFFFF')
    widths=[16,16,55,16,18,25,95,58,22,24]
    for j,v in enumerate(widths,1):s.column_dimensions[get_column_letter(j)].width=v
    for row in s.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=cell.column in (3,7))
    out=io.BytesIO();w.save(out);return out.getvalue()

def report_html(df,end,issue,order):
    E=lambda x:html.escape(str(x if pd.notna(x) else ''),quote=True)
    blocks=[]
    for region in order:
        part=df[df.region_code==region]
        if part.empty:continue
        blocks.append(f'<tr><td style="padding:18px 26px 6px"><div style="font:700 16px Arial;color:#111827">🌍 {E(region)}</div></td></tr>')
        for _,r in part.iterrows():
            tags=' '.join(f'<span style="display:inline-block;background:#e8f5e9;color:#2e7d32;border-radius:999px;padding:2px 8px;margin-right:4px">{E(t)}</span>' for t in re.split('[,，；;]',str(r.get('tags',''))) if t.strip())
            url=str(r.get('url',''));link=f'<a href="{E(url)}" style="display:inline-block;background:#2563eb;color:white;text-decoration:none;padding:8px 12px;border-radius:9px;font-size:12px" target="_blank">🔗 查看原文 ({E(r.get("url_label","原文"))})</a>' if url.startswith(('https://','http://')) else ''
            blocks.append(f'''<tr><td style="padding:10px 26px"><table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="border:1px solid #e5e7eb;border-radius:12px"><tr><td style="padding:14px"><div style="font:700 14px/1.5 Arial;color:#111827">{E(r.get('title',''))}</div><div style="font:12px/1.8 Arial;color:#6b7280;margin:8px 0">📍 {E(r.get('location',''))} | 📅 {E(r.get('date',''))} | {tags}</div><p style="font:13px/1.7 Arial;color:#374151;margin:8px 0 12px">{E(r.get('summary',''))}</p>{link}</td></tr></table></td></tr>''')
    return f'''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>全球汽车行业快讯</title></head><body style="margin:0;background:#f4f6f8"><table role="presentation" width="100%" cellpadding="0" cellspacing="0"><tr><td align="center" style="padding:24px 12px"><table role="presentation" width="680" cellpadding="0" cellspacing="0" style="max-width:680px;background:white;border-radius:16px"><tr><td style="padding:26px;background:#111827;color:white;text-align:center"><div style="font:700 28px Arial">全球汽车行业快讯</div><div style="font:14px Arial;color:#cbd5e1;margin-top:8px">📅 {end.year}年{end.month}月 | {E(issue)}</div></td></tr>{''.join(blocks)}<tr><td style="padding:18px 26px 26px;font:11px Arial;color:#9ca3af">最后更新：{date.today()}</td></tr></table></td></tr></table></body></html>'''

with st.sidebar:
    st.header('研究设置')
    start=st.date_input('开始日期',date.today()-timedelta(days=14))
    end=st.date_input('结束日期',date.today())
    issue=st.text_input('期数','第1期')
    file=st.file_uploader('可选：上传新闻 Excel 作为研究线索',type=['xlsx'])
    clues_text=st.text_area('可选：粘贴新闻标题或 URL（每行一条）',height=130)
    st.caption('不上传文件也能从零研究；上传内容不会直接当成事实。')
    with st.expander('API 设置（也可使用 Streamlit Secrets）'):
        sk=st.text_input('Serper API Key',value=secret('SERPER_API_KEY'),type='password')
        ak=st.text_input('OpenRouter API Key',value=secret('OPENROUTER_API_KEY'),type='password')
        model=st.text_input('OpenRouter Model',value=secret('OPENROUTER_MODEL','openai/gpt-4o-mini'))
    run=st.button('🔎 开始研究',type='primary',use_container_width=True)
    st.caption('规则可直接修改 searching_rules/config.json 和 RESEARCH_RULES.md')
if run:
    if start>end:st.error('开始日期不能晚于结束日期')
    elif not sk or not ak:st.error('请配置 Serper 和 OpenRouter API Key')
    else:
        try:
            clues,old=extract_clues(file,clues_text)
            bar=st.progress(0,text='准备搜索')
            rows,errors,n=research(start,end,clues,sk,ak,model,bar)
            st.session_state.rows=pd.DataFrame(rows)
            st.session_state.errors=errors
            st.session_state.found=n
            st.session_state.range=(start,end,issue)
            bar.empty()
        except Exception as e:st.error(f'研究未完成：{e}')
if 'rows' in st.session_state:
    df=st.session_state.rows
    st.success(f'检索到 {st.session_state.found} 个独立 URL；审核 {len(df)} 篇；自动精选 {int(df.selected.sum())} 篇。')
    if st.session_state.errors:
        with st.expander(f'搜索/分析异常 {len(st.session_state.errors)} 条'):
            st.code('\n'.join(st.session_state.errors[:80]))
    st.subheader('1 · 新闻研究与编辑')
    st.caption('可编辑标题、地区、日期、摘要、标签、链接、分数、纳入状态和排序。未核验新闻默认不选入。')
    edited=st.data_editor(df,hide_index=True,num_rows='dynamic',use_container_width=True,key='news_editor',column_config={'selected':st.column_config.CheckboxColumn('纳入快报'),'score':st.column_config.NumberColumn('分数',min_value=0,max_value=100),'region_code':st.column_config.SelectboxColumn('地区',options=REGIONS),'order':st.column_config.NumberColumn('顺序',min_value=0)},disabled=['source_date','source_access','source_domain'],height=520)
    st.session_state.rows=edited
    st.subheader('2 · 地区顺序')
    region_order=st.multiselect('按选择顺序显示地区',REGIONS,default=REGIONS)
    region_order=region_order+[x for x in REGIONS if x not in region_order]
    selected=edited[edited['selected'].fillna(False).astype(bool)].copy()
    selected['order']=pd.to_numeric(selected['order'],errors='coerce').fillna(9999)
    selected=selected.sort_values('order')
    selected['region_code']=selected['region_code'].where(selected['region_code'].isin(REGIONS),'亚太地区')
    selected=pd.concat([selected[selected.region_code==r] for r in region_order],ignore_index=True)
    s,e,iss=st.session_state.range
    xlsx=excel_bytes(selected,s,e,iss);page=report_html(selected,e,iss,region_order)
    st.subheader('3 · Outlook HTML 实时预览')
    st.components.v1.html(page,height=720,scrolling=True)
    a,b,c=st.columns(3)
    with a:st.download_button('📊 下载 Excel',xlsx,file_name=f'XSTAR_news_{s}_{e}.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',use_container_width=True)
    with b:st.download_button('📧 下载 HTML 快报',page,file_name=f'全球汽车行业快讯_{e}.html',mime='text/html',use_container_width=True)
    with c:st.download_button('🔍 下载全部研究审计 CSV',edited.to_csv(index=False).encode('utf-8-sig'),file_name=f'XSTAR_audit_{e}.csv',mime='text/csv',use_container_width=True)
else:
    st.info('输入日期即可从零搜索；Excel 和新闻列表均为可选。搜索需要真实 API Key。')
